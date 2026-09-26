"""Shared theme and slide builders for ICLUB bootcamp decks.

Every deck is built from these primitives so all tracks look identical.
Content lives in the per-track scripts; nothing in here is HTML-specific.
"""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# --------------------------------------------------------------------------
# Theme
# --------------------------------------------------------------------------

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

MARGIN = Inches(0.75)
CONTENT_W = Inches(11.833)

INK = RGBColor(0x0F, 0x17, 0x2A)
BODY = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
FAINT = RGBColor(0x94, 0xA3, 0xB8)

ACCENT = RGBColor(0x4F, 0x46, 0xE5)
ACCENT_SOFT = RGBColor(0xEE, 0xF2, 0xFF)

CODE_BG = RGBColor(0x0F, 0x17, 0x2A)
CODE_BG2 = RGBColor(0x1E, 0x29, 0x3B)
CODE_FG = RGBColor(0xE2, 0xE8, 0xF0)
CODE_KEY = RGBColor(0x93, 0xC5, 0xFD)
CODE_STR = RGBColor(0x86, 0xEF, 0xAC)
CODE_TXT = RGBColor(0xCB, 0xD5, 0xE1)

GOOD = RGBColor(0x05, 0x86, 0x53)
GOOD_SOFT = RGBColor(0xEC, 0xFD, 0xF3)
BAD = RGBColor(0xB9, 0x1C, 0x1C)
BAD_SOFT = RGBColor(0xFE, 0xF2, 0xF2)
WARN = RGBColor(0xB4, 0x53, 0x09)
WARN_SOFT = RGBColor(0xFF, 0xFB, 0xEB)

WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0xE2, 0xE8, 0xF0)

FONT = "Segoe UI"
FONT_BOLD = "Segoe UI Semibold"
FONT_MONO = "Consolas"


# --------------------------------------------------------------------------
# Primitives
# --------------------------------------------------------------------------


def new_deck():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _rect(slide, l, t, w, h, fill=None, radius=None, line=None, line_w=1):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, l, t, w, h)
    if radius:
        shp.adjustments[0] = radius
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def _text(
    slide,
    l,
    t,
    w,
    h,
    text,
    size=18,
    color=BODY,
    bold=False,
    font=FONT,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    line_spacing=1.25,
    space_after=0,
):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0

    lines = text.split("\n") if isinstance(text, str) else list(text)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        if space_after:
            p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = font
        run.font.color.rgb = color
    return box


def _bullets(slide, l, t, w, h, items, size=17, gap=9):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0

    for i, item in enumerate(items):
        text, level, kind = _parse_item(item)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.18
        p.space_after = Pt(gap if level == 0 else max(gap - 4, 2))
        p.level = level

        indent = Inches(0.0) if level == 0 else Inches(0.32)
        p.left_indent = indent
        first = p.runs[0] if p.runs else p.add_run()
        marker, color = _marker_for(kind)
        if marker:
            m = p.add_run()
            m.text = marker + "  "
            m.font.size = Pt(size if level == 0 else size - 2)
            m.font.bold = True
            m.font.name = FONT
            m.font.color.rgb = color

        r = p.add_run()
        r.text = text
        r.font.size = Pt(size if level == 0 else size - 2)
        r.font.name = FONT
        r.font.color.rgb = INK if kind != "plain" else BODY
        r.font.bold = kind in ("strong",)
    return box


def _parse_item(item):
    if isinstance(item, tuple):
        text, level, kind = item
        return text, level, kind
    if isinstance(item, str):
        return item, 0, "plain"
    raise TypeError(item)


def _marker_for(kind):
    return {
        "plain": ("\u2022", MUTED),
        "strong": ("\u25aa", ACCENT),
        "good": ("\u2713", GOOD),
        "bad": ("\u2715", BAD),
        "warn": ("!", WARN),
        "note": ("\u2192", ACCENT),
    }.get(kind, ("\u2022", MUTED))


# --------------------------------------------------------------------------
# Slide layouts
# --------------------------------------------------------------------------


def slide_title(prs, kicker, title, meta_lines, notes=""):
    s = _blank(prs)
    _rect(s, 0, 0, SLIDE_W, SLIDE_H, fill=INK)
    _rect(s, 0, Inches(4.72), Inches(1.1), Inches(0.055), fill=ACCENT)
    _text(s, MARGIN, Inches(1.9), Inches(11.0), Inches(0.4), kicker, 14, ACCENT, True, FONT_BOLD)
    _text(s, MARGIN, Inches(2.4), Inches(11.0), Inches(1.5), title, 46, WHITE, True, FONT_BOLD, line_spacing=1.0)
    _text(
        s,
        MARGIN,
        Inches(5.05),
        Inches(10.5),
        Inches(1.4),
        meta_lines,
        15,
        RGBColor(0x94, 0xA3, 0xB8),
        line_spacing=1.5,
    )
    _notes(s, notes)
    return s


def slide_divider(prs, number, title, duration, goal, topics, notes=""):
    s = _blank(prs)
    _rect(s, 0, 0, SLIDE_W, SLIDE_H, fill=INK)
    _text(s, MARGIN, Inches(1.55), Inches(3.0), Inches(0.4), f"SESSION {number}", 14, ACCENT, True, FONT_BOLD)
    _text(s, MARGIN, Inches(2.0), Inches(10.8), Inches(1.4), title, 38, WHITE, True, FONT_BOLD, line_spacing=1.05)
    _text(s, MARGIN, Inches(3.55), Inches(10.5), Inches(0.6), goal, 17, RGBColor(0xCB, 0xD5, 0xE1))
    _rect(s, MARGIN, Inches(4.35), Inches(1.1), Inches(0.045), fill=ACCENT)
    _text(s, MARGIN, Inches(4.75), Inches(11.0), Inches(1.3), topics, 14, RGBColor(0x94, 0xA3, 0xB8), line_spacing=1.45)
    _text(s, MARGIN, Inches(6.55), Inches(6.0), Inches(0.35), duration, 13, FAINT, font=FONT_BOLD)
    _notes(s, notes)
    return s


def slide_content(prs, title, items, kicker=None, notes="", size=17, gap=9):
    s = _blank(prs)
    _slide_head(s, title, kicker)
    _bullets(s, MARGIN, Inches(1.78), CONTENT_W, Inches(4.9), items, size=size, gap=gap)
    _notes(s, notes)
    return s


def slide_code(prs, title, code, caption=None, kicker=None, notes="", size=14, height=3.6):
    s = _blank(prs)
    _slide_head(s, title, kicker)
    top = Inches(1.78)
    _rect(s, MARGIN, top, CONTENT_W, Inches(height), fill=CODE_BG, radius=0.03)
    _text(
        s,
        MARGIN + Inches(0.35),
        top + Inches(0.28),
        CONTENT_W - Inches(0.7),
        Inches(height - 0.5),
        code,
        size,
        CODE_FG,
        font=FONT_MONO,
        line_spacing=1.35,
    )
    if caption:
        _text(
            s,
            MARGIN,
            top + Inches(height + 0.22),
            CONTENT_W,
            Inches(0.5),
            caption,
            14,
            MUTED,
        )
    _notes(s, notes)
    return s


def slide_two_col(prs, title, left_title, left_items, right_title, right_items,
                  kicker=None, notes="", accent_left=ACCENT, accent_right=ACCENT):
    s = _blank(prs)
    _slide_head(s, title, kicker)
    col_w = Inches(5.72)
    gap = Inches(0.4)
    top = Inches(1.8)
    for i, (head, its, acc) in enumerate(
        [(left_title, left_items, accent_left), (right_title, right_items, accent_right)]
    ):
        x = MARGIN + i * (col_w + gap)
        _rect(s, x, top, col_w, Inches(0.05), fill=acc)
        _text(s, x, top + Inches(0.22), col_w, Inches(0.4), head, 16, INK, True, FONT_BOLD)
        _bullets(s, x, top + Inches(0.78), col_w, Inches(4.0), its, size=15, gap=7)
    _notes(s, notes)
    return s


def slide_compare(prs, title, good_title, good_items, bad_title, bad_items,
                  kicker=None, notes=""):
    return slide_two_col(
        prs, title, good_title, good_items, bad_title, bad_items,
        kicker=kicker, notes=notes, accent_left=GOOD, accent_right=BAD,
    )


def slide_cards(prs, title, cards, kicker=None, notes="", cols=3, tone=ACCENT_SOFT):
    s = _blank(prs)
    _slide_head(s, title, kicker)
    rows = (len(cards) + cols - 1) // cols
    gap = Inches(0.3)
    w = int((CONTENT_W - gap * (cols - 1)) / cols)
    h = Inches(1.42)
    top = Inches(1.85)
    for i, (label, text) in enumerate(cards):
        r, c = divmod(i, cols)
        x = MARGIN + c * (w + gap)
        y = top + r * (h + Inches(0.26))
        _rect(s, x, y, w, h, fill=tone, radius=0.06)
        _rect(s, x, y, Inches(0.05), h, fill=ACCENT)
        _text(s, x + Inches(0.3), y + Inches(0.2), w - Inches(0.55), Inches(0.32), label, 14, INK, True, FONT_BOLD)
        _text(s, x + Inches(0.3), y + Inches(0.58), w - Inches(0.55), Inches(0.75), text, 12.5, BODY, line_spacing=1.2)
    _notes(s, notes)
    return s


def slide_table(prs, title, headers, rows, kicker=None, notes="", widths=None, size=13):
    s = _blank(prs)
    _slide_head(s, title, kicker)
    n = len(headers)
    total = 11.833
    widths = widths or [total / n] * n
    x = MARGIN
    top = Inches(1.85)
    rects = []
    for i, h in enumerate(headers):
        w = Inches(widths[i])
        _rect(s, x, top, w, Inches(0.46), fill=INK)
        _text(s, x + Inches(0.16), top + Inches(0.11), w - Inches(0.3), Inches(0.3), h, size, WHITE, True, FONT_BOLD)
        x += w
    y = top + Inches(0.46)
    for r, row in enumerate(rows):
        x = MARGIN
        h = Inches(0.52)
        _rect(s, MARGIN, y, Inches(total), h, fill=WHITE if r % 2 else RGBColor(0xF8, 0xFA, 0xFC))
        for i, cell in enumerate(row):
            w = Inches(widths[i])
            _text(s, x + Inches(0.16), y + Inches(0.14), w - Inches(0.3), Inches(0.3), str(cell), size, BODY)
            x += w
        y += h
    _rect(s, MARGIN, top, Inches(total), y - top, fill=None, line=BORDER, line_w=1)
    _notes(s, notes)
    return s


def slide_exercise(prs, label, title, tasks, hint=None, notes="", minutes=None):
    s = _blank(prs)
    _rect(s, 0, 0, Inches(0.16), SLIDE_H, fill=ACCENT)
    _text(s, MARGIN, Inches(0.72), Inches(6.0), Inches(0.35), label, 13, ACCENT, True, FONT_BOLD)
    _text(s, MARGIN, Inches(1.1), Inches(11.0), Inches(0.7), title, 30, INK, True, FONT_BOLD)
    if minutes:
        _text(s, Inches(10.6), Inches(0.75), Inches(1.95), Inches(0.35), minutes, 13, MUTED, align=PP_ALIGN.RIGHT)
    y = Inches(2.25)
    for i, t in enumerate(tasks):
        _rect(s, MARGIN, y + Inches(0.04), Inches(0.26), Inches(0.26), fill=ACCENT_SOFT, radius=0.5)
        _text(s, MARGIN, y + Inches(0.07), Inches(0.26), Inches(0.24), str(i + 1), 11, ACCENT, True, FONT_BOLD, align=PP_ALIGN.CENTER)
        _text(s, MARGIN + Inches(0.5), y, Inches(10.6), Inches(0.8), t, 17, BODY, line_spacing=1.3)
        y += Inches(0.78)
    if hint:
        _rect(s, MARGIN, y + Inches(0.15), CONTENT_W, Inches(0.95), fill=WARN_SOFT, radius=0.05)
        _rect(s, MARGIN, y + Inches(0.15), Inches(0.05), Inches(0.95), fill=WARN)
        _text(s, MARGIN + Inches(0.3), y + Inches(0.3), Inches(1.0), Inches(0.3), "HINT", 11, WARN, True, FONT_BOLD)
        _text(s, MARGIN + Inches(0.3), y + Inches(0.58), CONTENT_W - Inches(0.6), Inches(0.45), hint, 13.5, WARN)
    _notes(s, notes)
    return s


def slide_recap(prs, title, points, kicker=None, notes=""):
    s = _blank(prs)
    _slide_head(s, title, kicker)
    top = Inches(1.85)
    per_col = (len(points) + 1) // 2
    for col in range(2):
        chunk = points[col * per_col : (col + 1) * per_col]
        if not chunk:
            continue
        x = MARGIN + col * Inches(6.0)
        y = top
        for text in chunk:
            _rect(s, x, y + Inches(0.1), Inches(0.2), Inches(0.2), fill=ACCENT, radius=0.5)
            _text(s, x + Inches(0.42), y, Inches(5.3), Inches(0.75), text, 15, BODY, line_spacing=1.25)
            y += Inches(0.86)
    _notes(s, notes)
    return s


# --------------------------------------------------------------------------
# Chrome
# --------------------------------------------------------------------------


def _slide_head(s, title, kicker=None):
    if kicker:
        _text(s, MARGIN, Inches(0.52), Inches(9.0), Inches(0.3), kicker, 12.5, ACCENT, True, FONT_BOLD)
        y = Inches(0.86)
    else:
        y = Inches(0.66)
    _rect(s, MARGIN, y + Inches(0.06), Inches(0.07), Inches(0.52), fill=ACCENT)
    _text(s, MARGIN + Inches(0.24), y, Inches(11.4), Inches(0.7), title, 28, INK, True, FONT_BOLD, line_spacing=1.05)


def _notes(s, text):
    if text:
        s.notes_slide.notes_text_frame.text = text


def add_footer(prs, label):
    """Second pass: page numbers need the final slide count."""
    total = len(prs.slides._sldIdLst)
    for i, s in enumerate(prs.slides, start=1):
        _rect(s, 0, Inches(7.34), SLIDE_W, Inches(0.02), fill=BORDER)
        _text(s, MARGIN, Inches(6.96), Inches(9.0), Inches(0.28), label, 10, FAINT)
        _text(s, Inches(11.2), Inches(6.96), Inches(1.38), Inches(0.28), f"{i} / {total}", 10, FAINT, align=PP_ALIGN.RIGHT)


def save(prs, path):
    prs.save(path)
    return path
